"""Document upload: parsing PDF/DOCX/TXT, chunking, the admin API, and the database upgrade.

    cd backend && python -m unittest discover -s tests -p "test_documents*.py" -v
"""
import io
import unittest
import zipfile

import _env  # noqa: F401  (must be first: sets up the test environment)

from docx import Document as DocxDocument  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pypdf import PdfWriter  # noqa: E402
from sqlalchemy import create_engine, inspect, text  # noqa: E402

from app.database import Base, SessionLocal, upgrade_schema  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Document, KBArticle, KBEmbedding  # noqa: E402
from app.rag.semantic.index import semantic_index  # noqa: E402
from app.rag import ingestion as documents  # noqa: E402


# ---------------------------------------------------------------- tiny file builders
def make_docx(paragraphs, table=None) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                grid.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(lines) -> bytes:
    """A minimal one-page PDF with real (extractable) Helvetica text."""
    body = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(body)} >>\nstream\n{body}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for n, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def blank_pdf(password=None) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password:
        writer.encrypt(password, algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def bearer(token):
    return {"Authorization": f"Bearer {token}"}


RETURN_POLICY = [
    "سياسة استبدال الأجهزة",
    "يمكن استبدال الراوتر خلال 14 يومًا من تاريخ الاستلام بشرط أن يكون الجهاز بحالته الأصلية "
    "ومعه الفاتورة وجميع الملحقات.",
    "لا يشمل الاستبدال الأجهزة التي تعرضت لكسر أو سوائل. يتم فحص الجهاز في أقرب فرع قبل الموافقة على الاستبدال.",
]


# ---------------------------------------------------------------- pure functions
class ExtractionAndChunkingTests(unittest.TestCase):
    def test_docx_keeps_paragraphs_and_tables_in_order(self):
        data = make_docx(["العنوان", "فقرة أولى عن الراوتر."], table=[["الجهاز", "المدة"], ["راوتر", "14 يوم"]])
        out = documents.extract_text(data, "docx")
        self.assertEqual(out.split("\n"), ["العنوان", "فقرة أولى عن الراوتر.", "الجهاز | المدة", "راوتر | 14 يوم"])

    def test_pdf_text_is_extracted_and_lines_are_reflowed(self):
        data = make_pdf(["The router restarts", "automatically when power", "returns after an outage.",
                         "A second sentence follows here."])
        out = documents.extract_text(data, "pdf")
        self.assertIn("The router restarts automatically when power returns after an outage.", out)

    def test_pdf_hyphenated_word_is_rejoined(self):
        self.assertEqual(documents._reflow("the connec-\ntion is lost."), "the connection is lost.")

    def test_txt_encodings(self):
        sample = "سرعة الإنترنت بطيئة"
        self.assertEqual(documents.extract_text(sample.encode("utf-8-sig"), "txt"), sample)
        self.assertEqual(documents.extract_text(sample.encode("utf-16"), "txt"), sample)
        self.assertEqual(documents.extract_text(sample.encode("cp1256"), "txt"), sample)  # old Windows files
        with self.assertRaises(documents.DocumentError):
            documents.extract_text(b"\x00\x01\x02binary\x00", "txt")

    def test_arabic_presentation_forms_are_folded(self):
        self.assertEqual(documents.clean_text("\ufefb \u200bمرحبا\u200f"), "لا مرحبا")  # lam-alef ligature, invisibles
        self.assertEqual(documents.clean_text("a\r\n\r\n  b \t c "), "a\nb c")

    def test_chunks_are_small_overlap_and_lose_nothing(self):
        sentences = [f"الجملة رقم {i} تشرح خطوة من خطوات إعداد الراوتر في المنزل." for i in range(1, 61)]
        chunks = documents.chunk_text(" ".join(sentences))
        self.assertGreater(len(chunks), 4)
        for c in chunks:
            self.assertLessEqual(len(c.split()), documents.OVERLAP_MAX_WORDS + documents.MAX_WORDS)
        joined = " ".join(chunks)
        for s in sentences:
            self.assertIn(s, joined)
        # the closing sentence of a chunk opens the next one (context across the boundary)
        last = chunks[0].rsplit(". ", 1)[-1]
        self.assertTrue(chunks[1].startswith(last))

    def test_one_endless_sentence_is_cut_by_words(self):
        chunks = documents.chunk_text(" ".join(f"كلمة{i}" for i in range(400)))
        self.assertGreaterEqual(len(chunks), 4)
        for c in chunks:
            self.assertLessEqual(len(c.split()), documents.MAX_WORDS + documents.OVERLAP_MAX_WORDS)

    def test_junk_and_duplicates_are_dropped(self):
        self.assertEqual(documents.chunk_text("12"), [])
        self.assertEqual(documents.chunk_text("- 3 -\n\n  \n"), [])
        self.assertEqual(documents.chunk_text("Page 3 of 10"), [])
        text_ = "إعادة تشغيل الراوتر تحل أغلب مشاكل الاتصال بالإنترنت."
        self.assertEqual(documents.chunk_text(text_ + "\n" + text_), [text_ + "\n" + text_])  # one chunk, not two

    def test_file_info(self):
        self.assertEqual(documents._file_info("C:\\docs\\دليل_المستخدم  v2.PDF"), ("دليل المستخدم v2", "pdf"))
        with self.assertRaises(documents.DocumentError) as ctx:
            documents._file_info("virus.exe")
        self.assertEqual(ctx.exception.status_code, 415)


# ---------------------------------------------------------------- API
class DocumentApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._ctx = TestClient(app)
        cls.client = cls._ctx.__enter__()
        semantic_index.wait_idle(30)
        r = cls.client.post("/auth/login", json={"username": "superadmin", "password": "Sup3r-Admin-pass"})
        assert r.status_code == 200, r.text
        cls.admin = bearer(r.json()["access_token"])
        r = cls.client.post("/staff", headers=cls.admin, json={"username": "doc_agent", "password": "agent-pass-1"})
        assert r.status_code == 201, r.text
        r = cls.client.post("/auth/login", json={"username": "doc_agent", "password": "agent-pass-1"})
        cls.agent = bearer(r.json()["access_token"])
        r = cls.client.post("/customer-auth/signup", json={
            "username": "doc_customer", "password": "password123", "phone": "01099998888"})
        cls.cust = bearer(r.json()["access_token"])

    @classmethod
    def tearDownClass(cls):
        for d in cls.client.get("/admin/documents", headers=cls.admin).json():
            cls.client.delete(f"/admin/documents/{d['id']}", headers=cls.admin)
        semantic_index.wait_idle(30)
        cls._ctx.__exit__(None, None, None)

    def upload(self, name, data, product=None, headers=None):
        return self.client.post("/admin/documents", headers=headers or self.admin,
                                files={"file": (name, data)}, data={"product": product} if product else None)

    def ask(self, question):
        r = self.client.post("/tickets/message", headers=self.cust, json={"problem_text": question})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def doc_ids(self):
        return {d["id"] for d in self.client.get("/admin/documents", headers=self.admin).json()}

    # ---- the happy path
    def test_01_docx_upload_is_searchable_and_kept_out_of_the_kb_list(self):
        r = self.upload("سياسة_الاستبدال.docx", make_docx(RETURN_POLICY), product="سياسات")
        self.assertEqual(r.status_code, 201, r.text)
        doc = r.json()
        self.assertEqual((doc["title"], doc["file_type"], doc["product"], doc["uploaded_by"]),
                         ("سياسة الاستبدال", "docx", "سياسات", "superadmin"))
        self.assertGreaterEqual(doc["chunk_count"], 1)
        semantic_index.wait_idle(30)

        answer = self.ask("استبدال الراوتر")
        self.assertEqual((answer["source"], answer["status"]), ("kb", "resolved_by_rag"))
        self.assertIn("14 يومًا", answer["answer"])

        self.assertIn(doc["id"], self.doc_ids())
        kb = self.client.get("/admin/kb", headers=self.admin).json()
        self.assertFalse(any("استبدال" in (a["subject"] or "") for a in kb), "chunks must not flood the KB list")
        metrics = self.client.get("/admin/metrics", headers=self.admin).json()
        self.assertEqual(metrics["kb_article_count"], len(kb))
        self.assertGreaterEqual(metrics["document_count"], 1)

    def test_02_pdf_and_txt_upload(self):
        r = self.upload("warranty.pdf", make_pdf([
            "The router warranty lasts 24 months from the purchase date and covers manufacturing defects only."]))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["file_type"], "pdf")
        r = self.upload("install.txt", "يقوم الفني بتركيب الراوتر في موعد التركيب المحدد خلال 48 ساعة من الطلب.".encode())
        self.assertEqual(r.status_code, 201, r.text)
        semantic_index.wait_idle(30)
        self.assertEqual(self.ask("how long is the router warranty")["source"], "kb")
        self.assertIn("48 ساعة", self.ask("موعد التركيب بعد الطلب")["answer"])

    def test_03_same_file_twice_is_refused(self):
        data = make_docx(["دليل الدفع الإلكتروني للفواتير الشهرية عبر التطبيق والمحفظة."])
        self.assertEqual(self.upload("pay.docx", data).status_code, 201)
        again = self.upload("renamed-copy.docx", data)
        self.assertEqual(again.status_code, 409, again.text)

    # ---- bad input never creates anything
    def test_10_rejected_files(self):
        before = self.doc_ids()
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (documents.MAX_UNZIPPED_BYTES + 1))
        cases = [
            ("virus.exe", b"MZ....", 415),
            ("noext", b"hello world", 415),
            ("empty.txt", b"", 422),
            ("fake.pdf", b"this is not a pdf at all", 422),
            ("fake.docx", b"this is not a zip", 422),
            ("broken.docx", b"PK\x03\x04garbage-garbage-garbage", 422),
            ("binary.txt", b"\x00\x01\x02\x03" * 50, 422),
            ("scanned.pdf", blank_pdf(), 422),            # no text layer
            ("locked.pdf", blank_pdf("secret"), 422),     # password protected
            ("bomb.docx", bomb.getvalue(), 413),
            ("huge.txt", b"a " * (documents.MAX_UPLOAD_BYTES // 2 + 5), 413),
            ("noise.txt", b"1 2 3 4 5", 422),             # nothing searchable in it
        ]
        for name, data, expected in cases:
            r = self.upload(name, data)
            self.assertEqual(r.status_code, expected, f"{name}: {r.status_code} {r.text}")
        self.assertEqual(self.doc_ids(), before)

    # ---- permissions
    def test_20_permissions(self):
        self.assertEqual(self.client.get("/admin/documents").status_code, 401)
        self.assertEqual(self.client.get("/admin/documents", headers=self.cust).status_code, 401)
        self.assertEqual(self.client.get("/admin/documents", headers=self.agent).status_code, 200)  # read-only
        self.assertEqual(self.upload("a.txt", b"some searchable words here", headers=self.agent).status_code, 403)
        self.assertEqual(self.client.delete("/admin/documents/x", headers=self.agent).status_code, 403)
        self.assertEqual(self.client.post("/admin/documents", headers=self.admin).status_code, 422)  # no file field

    # ---- chunks belong to their document
    def test_30_chunks_cannot_be_edited_by_hand(self):
        with SessionLocal() as db:
            chunk = db.query(KBArticle).filter(KBArticle.document_id.isnot(None)).first()
            chunk_id = chunk.id
        body = {"product": "x", "ticket_type": "y", "subject": "z", "description": "", "resolution": "r"}
        self.assertEqual(self.client.put(f"/admin/kb/{chunk_id}", headers=self.admin, json=body).status_code, 409)
        self.assertEqual(self.client.delete(f"/admin/kb/{chunk_id}", headers=self.admin).status_code, 409)

    def test_31_reset_keeps_documents(self):
        before = self.doc_ids()
        self.assertTrue(before)
        self.assertEqual(self.client.post("/admin/reset", headers=self.admin).status_code, 200)
        self.assertEqual(self.doc_ids(), before)

    def test_90_delete_removes_chunks_vectors_and_answers(self):
        r = self.upload("unique.txt", "جدول تدريب الموظفين الجدد يبدأ كل أحد ويستمر ثلاثة أسابيع في المقر الرئيسي.".encode())
        self.assertEqual(r.status_code, 201, r.text)
        doc_id = r.json()["id"]
        semantic_index.wait_idle(30)
        self.assertIn("تدريب", self.ask("تدريب الموظفين الجدد المقر الرئيسي")["answer"])

        self.assertEqual(self.client.delete(f"/admin/documents/{doc_id}", headers=self.admin).status_code, 200)
        semantic_index.wait_idle(30)
        with SessionLocal() as db:
            self.assertEqual(db.query(KBArticle).filter(KBArticle.document_id == doc_id).count(), 0)
            self.assertIsNone(db.get(Document, doc_id))
            orphans = db.query(KBEmbedding).filter(~KBEmbedding.article_id.in_(db.query(KBArticle.id))).count()
            self.assertEqual(orphans, 0)
        self.assertNotIn("تدريب", self.ask("تدريب الموظفين الجدد المقر الرئيسي")["answer"])
        self.assertEqual(self.client.delete(f"/admin/documents/{doc_id}", headers=self.admin).status_code, 404)

    def test_91_large_document_becomes_many_chunks(self):
        paragraphs = [f"الفقرة {i}: " + " ".join(f"خطوة{i}_{j} في تشغيل الراوتر" for j in range(12)) + "." for i in range(40)]
        r = self.upload("manual.docx", make_docx(paragraphs))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertGreater(r.json()["chunk_count"], 10)
        with SessionLocal() as db:
            n = db.query(KBArticle).filter(KBArticle.document_id == r.json()["id"]).count()
        self.assertEqual(n, r.json()["chunk_count"])


# ---------------------------------------------------------------- upgrading an existing database
class UpgradeSchemaTests(unittest.TestCase):
    def test_old_database_gets_document_id_without_losing_data(self):
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine, tables=[Document.__table__])
        with engine.begin() as conn:  # the knowledge-base table exactly as it was before documents existed
            conn.execute(text(
                "CREATE TABLE kb_articles (id INTEGER PRIMARY KEY AUTOINCREMENT, product VARCHAR(200), "
                "ticket_type VARCHAR(200), subject VARCHAR(500), description TEXT, resolution TEXT, status VARCHAR(50), "
                "priority VARCHAR(20), channel VARCHAR(50), created_at DATETIME, updated_at DATETIME)"))
            conn.execute(text("INSERT INTO kb_articles (product, subject, resolution) VALUES ('p', 'old article', 'r')"))
        upgrade_schema(engine)
        upgrade_schema(engine)  # running it again is harmless
        columns = {c["name"] for c in inspect(engine).get_columns("kb_articles")}
        self.assertIn("document_id", columns)
        with engine.connect() as conn:
            row = conn.execute(text("SELECT subject, document_id FROM kb_articles")).one()
        self.assertEqual(tuple(row), ("old article", None))


if __name__ == "__main__":
    unittest.main()
